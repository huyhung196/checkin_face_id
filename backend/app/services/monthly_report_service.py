import io
import calendar
from datetime import datetime, date, timedelta
from typing import Dict, List, Any, Optional, Tuple

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.config import get_vietnam_now, DB_FILE
from app.database import get_db

# Mã màu thương hiệu chuẩn theo file report_example.xlsx
COLOR_HEADER_BG = "1155CC"       # Xanh dương đậm Header
COLOR_HEADER_TEXT = "FFFFFF"     # Chữ trắng
COLOR_ZEBRA_BG = "CFE2F3"        # Xanh nhạt xen kẽ (Zebra stripe)
COLOR_TITLE_TEXT = "C55A11"      # Cam đậm tiêu đề Báo Cáo Chấm Công MEBIECO
COLOR_EMP_TEXT = "0B5394"        # Xanh đậm tên nhân viên
COLOR_BORDER = "B7B7B7"          # Xám viền mỏng
COLOR_ALERT_RED = "C00000"       # Đỏ cảnh báo vi phạm
COLOR_PERM_GREEN = "059669"      # Xanh lá có phép
COLOR_MUTED_TEXT = "595959"      # Xám ghi chú

VIETNAMESE_WEEKDAYS = {
    0: "Thứ 2",
    1: "Thứ 3",
    2: "Thứ 4",
    3: "Thứ 5",
    4: "Thứ 6",
    5: "Thứ 7",
    6: "Chủ Nhật"
}


def _get_border():
    thin = Side(border_style="thin", color=COLOR_BORDER)
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def get_monthly_attendance_data(
    year_month: str = "",
    from_date: str = "",
    to_date: str = "",
    employee_id: Optional[int] = None
) -> Dict[str, Any]:
    """
    Thu thập và gộp thông minh các lượt Check In - Check Out của từng nhân viên theo ngày.
    Mặc định: Lấy từ ngày đầu tháng đến ngày hiện tại nếu là tháng này.
    """
    now = get_vietnam_now()

    if not year_month:
        year_month = now.strftime("%Y-%m")

    try:
        yr, mo = map(int, year_month.split("-"))
    except Exception:
        yr, mo = now.year, now.month
        year_month = f"{yr:04d}-{mo:02d}"

    _, last_day_of_month = calendar.monthrange(yr, mo)

    # Xác định khoảng ngày
    if not from_date:
        from_date = f"{year_month}-01"

    if not to_date:
        if yr == now.year and mo == now.month:
            to_date = now.strftime("%Y-%m-%d")
        else:
            to_date = f"{year_month}-{last_day_of_month:02d}"

    conn = get_db()
    cursor = conn.cursor()

    # 1. Lấy danh sách nhân viên hiện tại
    cursor.execute("SELECT id, employee_code, full_name, department, position FROM employees ORDER BY employee_code ASC, full_name ASC")
    employees = [dict(row) for row in cursor.fetchall()]

    emp_by_id = {e["id"]: e for e in employees}
    emp_by_code = {e["employee_code"].lower().strip(): e for e in employees if e.get("employee_code")}
    emp_by_name = {e["full_name"].lower().strip(): e for e in employees if e.get("full_name")}

    selected_emp = emp_by_id.get(employee_id) if employee_id else None

    # 2. Lấy logs điểm danh trong khoảng ngày
    start_ts = f"{from_date}T00:00:00"
    end_ts = f"{to_date}T23:59:59.999999"

    log_query = """
        SELECT id, timestamp, formatted_time, employee_id, employee_code, user_name,
               check_type, working_hours, working_duration,
               attendance_status, late_minutes, early_minutes,
               has_permission, permission_note
        FROM checkin_logs
        WHERE timestamp >= ? AND timestamp <= ?
        ORDER BY timestamp ASC, id ASC
    """
    cursor.execute(log_query, (start_ts, end_ts))
    all_logs = [dict(row) for row in cursor.fetchall()]
    conn.close()

    # Nhận diện nhân viên cho từng log (Khớp thông minh: ID -> Mã NV -> Họ Tên)
    def match_employee_for_log(log_item):
        l_eid = log_item.get("employee_id")
        l_code = (log_item.get("employee_code") or "").strip()
        l_name = (log_item.get("user_name") or "").strip()

        if l_eid and l_eid in emp_by_id:
            return emp_by_id[l_eid]
        if l_code and l_code.lower() in emp_by_code:
            return emp_by_code[l_code.lower()]
        if l_name and l_name.lower() in emp_by_name:
            return emp_by_name[l_name.lower()]
        return {
            "id": l_eid or f"log_{l_name}",
            "employee_code": l_code or "---",
            "full_name": l_name or "Nhân viên",
            "department": "---",
            "position": "---"
        }

    # Gom log theo nhân viên đã đối soát -> date_str (YYYY-MM-DD)
    canonical_emp_map: Dict[Any, Dict[str, Any]] = {}
    for emp in employees:
        canonical_emp_map[emp["id"]] = {
            "emp": emp,
            "days": {}
        }

    for log in all_logs:
        m_emp = match_employee_for_log(log)
        m_id = m_emp["id"]

        # Nếu đang lọc theo 1 nhân viên cụ thể
        if selected_emp:
            is_same_id = (m_id == selected_emp["id"])
            is_same_name = (m_emp.get("full_name", "").lower().strip() == selected_emp.get("full_name", "").lower().strip())
            if not (is_same_id or is_same_name):
                continue

        if m_id not in canonical_emp_map:
            canonical_emp_map[m_id] = {
                "emp": m_emp,
                "days": {}
            }

        ts = log.get("timestamp", "")
        if not ts:
            continue
        d_str = ts[:10]  # YYYY-MM-DD
        if d_str not in canonical_emp_map[m_id]["days"]:
            canonical_emp_map[m_id]["days"][d_str] = []
        canonical_emp_map[m_id]["days"][d_str].append(log)

    # 3. Tiến hành gộp Check In & Check Out theo từng ngày của từng nhân viên
    report_rows = []
    emp_summaries = []

    for m_id, emp_container in canonical_emp_map.items():
        emp = emp_container["emp"]
        user_days = emp_container["days"]
        emp_name = emp.get("full_name", "")
        emp_code = emp.get("employee_code", "") or "---"
        dept = emp.get("department", "") or "---"

        # Lấy danh sách các ngày có chấm công và sắp xếp
        sorted_dates = sorted(user_days.keys())

        # Thống kê tổng hợp cho nhân viên này
        total_attendance_days = len(sorted_dates)
        total_work_credit = 0.0
        late_count = 0
        total_late_mins = 0
        early_count = 0
        total_early_mins = 0
        no_checkout_count = 0
        permission_count = 0

        emp_detail_rows = []

        for d_str in sorted_dates:
            day_logs = user_days[d_str]
            if not day_logs:
                continue

            dt_obj = datetime.strptime(d_str, "%Y-%m-%d")
            day_num = dt_obj.day
            weekday_str = VIETNAMESE_WEEKDAYS.get(dt_obj.weekday(), "")

            # Phân tích các log trong ngày
            # Quy tắc: Log sớm nhất là Check In, Log muộn nhất là Check Out
            first_log = day_logs[0]
            last_log = day_logs[-1] if len(day_logs) > 1 else None

            check_in_time = None
            check_out_time = None
            working_dur = "---"
            check_status_items = []
            perm_notes = []
            day_work_credit = 1.0

            if last_log is None:
                # Chỉ có đúng 1 lượt quét trong ngày
                log_type = first_log.get("check_type", "Vào Ca")
                log_time_str = first_log.get("timestamp", "").split("T")[1][:5] if "T" in first_log.get("timestamp", "") else ""

                if log_type == "Vào Ca":
                    check_in_time = log_time_str
                    check_out_time = ""
                    check_status_items.append("Không Check Out")
                    no_checkout_count += 1
                    day_work_credit = 0.5  # Thiếu check out tính 0.5 công

                    # Kiểm tra vào trễ
                    if first_log.get("attendance_status") == "Đi Trễ":
                        late_m = first_log.get("late_minutes", 0) or 0
                        late_count += 1
                        total_late_mins += late_m
                        if first_log.get("has_permission"):
                            check_status_items.append("Đi Trễ - Có Phép")
                            permission_count += 1
                        else:
                            check_status_items.append(f"Đi Trễ ({late_m}p)" if late_m else "Đi Trễ")
                        if first_log.get("permission_note"):
                            perm_notes.append(first_log["permission_note"])
                else:
                    # Là Tan Ca mà không có Vào Ca
                    check_in_time = ""
                    check_out_time = log_time_str
                    check_status_items.append("Không Check In")
                    day_work_credit = 0.5
                    if first_log.get("attendance_status") == "Về Sớm":
                        early_m = first_log.get("early_minutes", 0) or 0
                        early_count += 1
                        total_early_mins += early_m
                        if first_log.get("has_permission"):
                            check_status_items.append("Về Sớm - Có Phép")
                            permission_count += 1
                        else:
                            check_status_items.append(f"Về Sớm ({early_m}p)" if early_m else "Về Sớm")
            else:
                # Có từ 2 lượt quét trở lên: Log đầu là Vào Ca, Log cuối là Tan Ca
                check_in_time = first_log.get("timestamp", "").split("T")[1][:5] if "T" in first_log.get("timestamp", "") else ""
                check_out_time = last_log.get("timestamp", "").split("T")[1][:5] if "T" in last_log.get("timestamp", "") else ""

                try:
                    from app.services.shift_service import calculate_working_duration
                    f_dt = datetime.fromisoformat(first_log["timestamp"])
                    l_dt = datetime.fromisoformat(last_log["timestamp"])
                    dur_calc = calculate_working_duration(f_dt, l_dt)
                    working_dur = dur_calc["working_duration"]
                    net_h = dur_calc["working_hours"]
                    if net_h >= 7.0:
                        day_work_credit = 1.0
                    elif net_h >= 3.5:
                        day_work_credit = 0.5
                    else:
                        day_work_credit = 0.5 if not (last_log.get("attendance_status") == "Về Sớm") else 0.0
                except Exception:
                    working_dur = last_log.get("working_duration") or "---"
                    day_work_credit = 1.0

                # Kiểm tra vào ca (Trễ)
                is_late = first_log.get("attendance_status") == "Đi Trễ"
                late_m = first_log.get("late_minutes", 0) or 0
                has_late_perm = bool(first_log.get("has_permission"))

                # Kiểm tra tan ca (Sớm)
                is_early = last_log.get("attendance_status") == "Về Sớm"
                early_m = last_log.get("early_minutes", 0) or 0
                has_early_perm = bool(last_log.get("has_permission"))

                if is_late:
                    late_count += 1
                    total_late_mins += late_m
                if is_early:
                    early_count += 1
                    total_early_mins += early_m

                if has_late_perm or has_early_perm:
                    permission_count += 1

                if first_log.get("permission_note"):
                    perm_notes.append(first_log["permission_note"])
                if last_log.get("permission_note") and last_log.get("permission_note") not in perm_notes:
                    perm_notes.append(last_log["permission_note"])

                # Định dạng nhãn Cột Check chuẩn theo phong cách report_example.xlsx:
                if is_late and is_early:
                    if has_late_perm and has_early_perm:
                        check_status_items.append("Đi Trễ, Về Sớm - Có Phép")
                    elif has_late_perm:
                        check_status_items.append("Đi Trễ (Có Phép), Về Sớm")
                    elif has_early_perm:
                        check_status_items.append("Đi Trễ, Về Sớm (Có Phép)")
                    else:
                        check_status_items.append("Đi Trễ, Về Sớm")
                elif is_late:
                    if has_late_perm:
                        check_status_items.append("Đi Trễ - Có Phép")
                    else:
                        check_status_items.append(f"Đi Trễ ({late_m}p)" if late_m > 0 else "Đi Trễ")
                elif is_early:
                    if has_early_perm:
                        check_status_items.append("Về Sớm - Có Phép")
                    else:
                        check_status_items.append(f"Về Sớm ({early_m}p)" if early_m > 0 else "Về Sớm")
                else:
                    # Đúng giờ
                    check_status_items = []

            # Ghép chuỗi Check status
            check_label = ", ".join(check_status_items) if check_status_items else ""
            permission_note_str = " | ".join(perm_notes) if perm_notes else ""

            total_work_credit += day_work_credit

            emp_detail_rows.append({
                "employee_code": emp_code,
                "user_name": emp_name,
                "department": dept,
                "day": day_num,
                "date_str": d_str,
                "weekday": weekday_str,
                "check_in": check_in_time or "",
                "check_out": check_out_time or "",
                "working_duration": working_dur,
                "check_status": check_label,
                "permission_note": permission_note_str,
                "work_credit": day_work_credit
            })

        emp_summaries.append({
            "employee_code": emp_code,
            "full_name": emp_name,
            "department": dept,
            "total_days": total_attendance_days,
            "total_credits": round(total_work_credit, 1),
            "late_count": late_count,
            "total_late_mins": total_late_mins,
            "early_count": early_count,
            "total_early_mins": total_early_mins,
            "no_checkout_count": no_checkout_count,
            "permission_count": permission_count
        })

        report_rows.append({
            "employee": emp,
            "rows": emp_detail_rows
        })

    return {
        "year_month": year_month,
        "from_date": from_date,
        "to_date": to_date,
        "company_name": "MEBIECO",
        "report_data": report_rows,
        "summaries": emp_summaries
    }


def generate_monthly_attendance_excel(
    year_month: str = "",
    from_date: str = "",
    to_date: str = "",
    employee_id: Optional[int] = None
) -> Tuple[io.BytesIO, str]:
    """
    Tạo file Excel (.xlsx) chuẩn phong cách file report_example.xlsx:
    - Tiêu đề Cam đậm MEBIECO
    - Header xanh đậm #1155CC
    - Merge ô tên Nhân Viên
    - Dòng xen kẽ trắng và xanh nhạt #CFE2F3
    - Tự động gộp Check In và Check Out của 1 nhân viên thành 1 hàng
    - Bổ sung Sheet 2: Bảng Tổng Hợp Công Tháng cho HR
    """
    data = get_monthly_attendance_data(
        year_month=year_month,
        from_date=from_date,
        to_date=to_date,
        employee_id=employee_id
    )

    wb = openpyxl.Workbook()

    # Font definitions
    font_main_title = Font(name="Roboto", size=15, bold=True, color=COLOR_TITLE_TEXT)
    font_subtitle = Font(name="Roboto", size=11, bold=False, color=COLOR_TITLE_TEXT)
    font_header = Font(name="Roboto", size=11, bold=True, color=COLOR_HEADER_TEXT)
    font_emp_name = Font(name="Roboto", size=11, bold=True, color=COLOR_EMP_TEXT)
    font_cell = Font(name="Roboto", size=10, bold=False, color="000000")
    font_cell_bold = Font(name="Roboto", size=10, bold=True, color="000000")
    font_alert = Font(name="Roboto", size=10, bold=True, color=COLOR_ALERT_RED)
    font_perm = Font(name="Roboto", size=10, bold=True, color=COLOR_PERM_GREEN)
    font_note = Font(name="Roboto", size=9, italic=True, color=COLOR_MUTED_TEXT)

    # Fill definitions
    fill_header = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_zebra = PatternFill(start_color=COLOR_ZEBRA_BG, end_color=COLOR_ZEBRA_BG, fill_type="solid")

    border_thin = _get_border()

    # -------------------------------------------------------------
    # SHEET 1: Chi Tiết Chấm Công (Style theo report_example.xlsx)
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Báo Cáo Chấm Công"
    ws1.views.sheetView[0].showGridLines = True

    # Cấu trúc cột Sheet 1
    # Col 1: Nhân Viên
    # Col 2: Ngày
    # Col 3: Thứ
    # Col 4: Check In
    # Col 5: Check Out
    # Col 6: Thời Lượng
    # Col 7: Check (Tình Trạng)
    # Col 8: Ghi Chú Phép
    # Col 9: Công
    total_cols = 9

    # Dòng 1-2: Title
    ws1.merge_cells(start_row=1, start_column=1, end_row=2, end_column=total_cols)
    c_title = ws1.cell(1, 1, value="Báo Cáo Chấm Công MEBIECO")
    c_title.font = font_main_title
    c_title.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 22
    ws1.row_dimensions[2].height = 22

    # Dòng 3: Subtitle
    yr_str, mo_str = data["year_month"].split("-")
    ws1.merge_cells(start_row=3, start_column=1, end_row=3, end_column=total_cols)
    from_d_fmt = datetime.strptime(data["from_date"], "%Y-%m-%d").strftime("%d/%m/%Y")
    to_d_fmt = datetime.strptime(data["to_date"], "%Y-%m-%d").strftime("%d/%m/%Y")
    c_sub = ws1.cell(3, 1, value=f"Tháng {int(mo_str)}/{yr_str} (Từ {from_d_fmt} đến {to_d_fmt})")
    c_sub.font = font_subtitle
    c_sub.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[3].height = 20

    # Dòng 4: Headers (Row 4)
    headers_sheet1 = [
        "Nhân Viên",
        "Ngày",
        "Thứ",
        "Check In",
        "Check Out",
        "Thời Lượng",
        "Check",
        "Ghi Chú / Lý Do Phép",
        "Công"
    ]
    ws1.row_dimensions[4].height = 28
    for col_idx, h_text in enumerate(headers_sheet1, 1):
        c = ws1.cell(4, col_idx, value=h_text)
        c.font = font_header
        c.fill = fill_header
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border_thin

    current_row = 5

    for emp_group in data["report_data"]:
        emp_info = emp_group["employee"]
        rows = emp_group["rows"]

        if not rows:
            # Nhân viên chưa có log nào trong khoảng ngày này
            continue

        emp_display = f"{emp_info.get('full_name', '')}\n({emp_info.get('employee_code', '')})" if emp_info.get('employee_code') else emp_info.get('full_name', '')
        start_emp_row = current_row

        for idx, r in enumerate(rows):
            # Xen kẽ màu dòng (chẵn/lẻ) giống mẫu
            row_fill = fill_zebra if idx % 2 == 1 else fill_white

            ws1.row_dimensions[current_row].height = 24

            # Col 1: Nhân Viên (Tạm gán, tí merge)
            c1 = ws1.cell(current_row, 1, value=emp_display)
            c1.font = font_emp_name
            c1.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c1.border = border_thin

            # Col 2: Ngày
            c2 = ws1.cell(current_row, 2, value=r["day"])
            c2.font = font_cell_bold
            c2.fill = row_fill
            c2.alignment = Alignment(horizontal="center", vertical="center")
            c2.border = border_thin

            # Col 3: Thứ
            c3 = ws1.cell(current_row, 3, value=r["weekday"])
            c3.font = font_cell
            c3.fill = row_fill
            c3.alignment = Alignment(horizontal="center", vertical="center")
            c3.border = border_thin

            # Col 4: Check In
            c4 = ws1.cell(current_row, 4, value=r["check_in"] or "---")
            c4.font = font_cell
            c4.fill = row_fill
            c4.alignment = Alignment(horizontal="center", vertical="center")
            c4.border = border_thin

            # Col 5: Check Out
            c5 = ws1.cell(current_row, 5, value=r["check_out"] or "---")
            c5.font = font_cell
            c5.fill = row_fill
            c5.alignment = Alignment(horizontal="center", vertical="center")
            c5.border = border_thin

            # Col 6: Thời Lượng
            c6 = ws1.cell(current_row, 6, value=r["working_duration"] if r["working_duration"] != "---" else "")
            c6.font = font_cell
            c6.fill = row_fill
            c6.alignment = Alignment(horizontal="center", vertical="center")
            c6.border = border_thin

            # Col 7: Check (Tình trạng)
            check_val = r["check_status"]
            c7 = ws1.cell(current_row, 7, value=check_val or None)
            if "Không Check" in check_val or "Đi Trễ" in check_val or "Về Sớm" in check_val:
                c7.font = font_perm if "Có Phép" in check_val else font_alert
            else:
                c7.font = font_cell
            c7.fill = row_fill
            c7.alignment = Alignment(horizontal="left", vertical="center")
            c7.border = border_thin

            # Col 8: Ghi chú phép
            c8 = ws1.cell(current_row, 8, value=r["permission_note"] or None)
            c8.font = font_note
            c8.fill = row_fill
            c8.alignment = Alignment(horizontal="left", vertical="center")
            c8.border = border_thin

            # Col 9: Công
            c9 = ws1.cell(current_row, 9, value=r["work_credit"])
            c9.font = font_cell_bold
            c9.fill = row_fill
            c9.alignment = Alignment(horizontal="center", vertical="center")
            c9.border = border_thin

            current_row += 1

        # Gộp ô tên Nhân Viên cho tất cả các dòng của nhân viên đó (Merge cell)
        end_emp_row = current_row - 1
        if end_emp_row > start_emp_row:
            ws1.merge_cells(start_row=start_emp_row, start_column=1, end_row=end_emp_row, end_column=1)
            # Viền cho các ô đã gộp
            for r_idx in range(start_emp_row, end_emp_row + 1):
                ws1.cell(r_idx, 1).border = border_thin

    # Tự động căn chỉnh độ rộng cột Sheet 1
    col_widths_sheet1 = {
        1: 22,  # Nhân Viên
        2: 8,   # Ngày
        3: 11,  # Thứ
        4: 12,  # Check In
        5: 12,  # Check Out
        6: 14,  # Thời Lượng
        7: 26,  # Check
        8: 26,  # Ghi Chú Phép
        9: 10   # Công
    }
    for col_idx, width in col_widths_sheet1.items():
        ws1.column_dimensions[get_column_letter(col_idx)].width = width

    # -------------------------------------------------------------
    # SHEET 2: Bảng Tổng Hợp Công HR (Phục vụ chốt lương cuối tháng)
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Tổng Hợp Công HR")
    ws2.views.sheetView[0].showGridLines = True

    # Title Sheet 2
    ws2.merge_cells("A1:K2")
    c_title2 = ws2.cell(1, 1, value=f"BẢNG TỔNG HỢP CÔNG CHẤM CÔNG THÁNG {int(mo_str)}/{yr_str}")
    c_title2.font = font_main_title
    c_title2.alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 22
    ws2.row_dimensions[2].height = 22

    # Subtitle Sheet 2
    ws2.merge_cells("A3:K3")
    c_sub2 = ws2.cell(3, 1, value=f"Thời gian: Từ {from_d_fmt} đến {to_d_fmt} | Đơn vị: MEBIECO")
    c_sub2.font = font_subtitle
    c_sub2.alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[3].height = 20

    headers_sheet2 = [
        "STT",
        "Mã NV",
        "Họ và Tên",
        "Phòng Ban",
        "Số Ngày Điểm Danh",
        "Tổng Công",
        "Số Lần Trễ",
        "Tổng Phút Trễ",
        "Số Lần Về Sớm",
        "Tổng Phút Về Sớm",
        "Quên Check Out"
    ]
    ws2.row_dimensions[4].height = 28
    for col_idx, h_text in enumerate(headers_sheet2, 1):
        c = ws2.cell(4, col_idx, value=h_text)
        c.font = font_header
        c.fill = fill_header
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border_thin

    s2_row = 5
    for idx, s in enumerate(data["summaries"], 1):
        row_fill = fill_zebra if idx % 2 == 0 else fill_white
        ws2.row_dimensions[s2_row].height = 22

        vals = [
            idx,
            s["employee_code"],
            s["full_name"],
            s["department"],
            s["total_days"],
            s["total_credits"],
            s["late_count"],
            s["total_late_mins"],
            s["early_count"],
            s["total_early_mins"],
            s["no_checkout_count"]
        ]

        for col_idx, val in enumerate(vals, 1):
            c = ws2.cell(s2_row, col_idx, value=val)
            c.font = font_cell_bold if col_idx in [2, 3, 6] else font_cell
            c.fill = row_fill
            c.border = border_thin
            if col_idx in [1, 2, 5, 6, 7, 8, 9, 10, 11]:
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")

        s2_row += 1

    # Dòng Tổng Cộng
    if data["summaries"]:
        ws2.row_dimensions[s2_row].height = 24
        c_tot_label = ws2.cell(s2_row, 3, value="TỔNG CỘNG")
        c_tot_label.font = Font(name="Roboto", size=10, bold=True, color="000000")
        c_tot_label.alignment = Alignment(horizontal="center", vertical="center")

        # SUM công thức
        for col_idx in [5, 6, 7, 8, 9, 10, 11]:
            col_letter = get_column_letter(col_idx)
            c_sum = ws2.cell(s2_row, col_idx, value=f"=SUM({col_letter}5:{col_letter}{s2_row-1})")
            c_sum.font = Font(name="Roboto", size=10, bold=True, color=COLOR_EMP_TEXT)
            c_sum.alignment = Alignment(horizontal="center", vertical="center")
            c_sum.border = border_thin

        for c_i in range(1, 12):
            ws2.cell(s2_row, c_i).border = border_thin

    col_widths_sheet2 = {
        1: 6,   # STT
        2: 12,  # Mã NV
        3: 24,  # Họ và Tên
        4: 18,  # Phòng Ban
        5: 18,  # Số Ngày Điểm Danh
        6: 14,  # Tổng Công
        7: 13,  # Số Lần Trễ
        8: 15,  # Tổng Phút Trễ
        9: 15,  # Số Lần Về Sớm
        10: 17, # Tổng Phút Về Sớm
        11: 16  # Quên Check Out
    }
    for col_idx, width in col_widths_sheet2.items():
        ws2.column_dimensions[get_column_letter(col_idx)].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Bao_Cao_Cham_Cong_Thang_{int(mo_str):02d}_{yr_str}.xlsx"
    return output, filename
